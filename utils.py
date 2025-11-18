import datetime
import os
import openpyxl
import logging
from telebot import types
from config import bot
from date_utils import is_valid_date, is_weekend

logger = logging.getLogger(__name__)

def ask_date(message, fio, phone_number, massage_type, next_step_handler, old_date=None, old_time=None):
    """Показывает клавиатуру с датами для записи"""
    keyboard = types.ReplyKeyboardMarkup(resize_keyboard=True)
    today = datetime.date.today()
    for i in range(7):
        date = today + datetime.timedelta(days=i)
        if not is_weekend(date):
            keyboard.add(date.strftime("%d.%m.%Y"))
    sent = bot.send_message(message.chat.id, "Выберите дату:", reply_markup=keyboard)
    # Регистрируем следующий шаг по chat_id, чтобы обработка сработала вне зависимости от объекта message
    bot.register_next_step_handler_by_chat_id(message.chat.id, ask_time, fio, phone_number, massage_type, next_step_handler, old_date, old_time)

def ask_time(message, fio, phone_number, massage_type, next_step_handler, date, old_date=None, old_time=None):
    """Показывает клавиатуру с доступным временем для выбранной даты"""
    date = (message.text or "").strip()
    if not is_valid_date(date):
        bot.send_message(message.chat.id, "Неверная дата. Выберите дату из предложенных.", reply_markup=types.ReplyKeyboardRemove())
        # Регистрируем по chat_id, чтобы следующий ввод корректно попал в обработчик
        bot.register_next_step_handler_by_chat_id(message.chat.id, ask_date, fio, phone_number, massage_type, next_step_handler, old_date, old_time)
        return

    available_times = get_available_times(date)
    if available_times:
        keyboard = types.ReplyKeyboardMarkup(resize_keyboard=True)
        for time in available_times:
            keyboard.add(time)
        sent = bot.send_message(message.chat.id, "Выберите время:", reply_markup=keyboard)
        # Регистрируем следующий шаг по chat_id — далее будет вызван next_step_handler
        bot.register_next_step_handler_by_chat_id(message.chat.id, next_step_handler, fio, phone_number, massage_type, date, old_date, old_time)
    else:
        bot.send_message(message.chat.id, "Нет свободного времени. Выберите другой день.", reply_markup=types.ReplyKeyboardRemove())
        bot.register_next_step_handler_by_chat_id(message.chat.id, ask_date, fio, phone_number, massage_type, next_step_handler, old_date, old_time)

import threading
excel_lock = threading.RLock()

def get_available_times(date):
    """Возвращает список доступного времени для записи на выбранную дату"""
    fixed_times = [f"{hour:02d}:{minute:02d}" for hour in range(9, 18) for minute in (0, 30)]
    booked_times = []
    try:
        with excel_lock:
            logger.info(f"Чтение db.xlsx для поиска занятых времён на дату={date}")
            wb = openpyxl.load_workbook('db.xlsx')
            sheet = wb.active
            booked_times = []
            for row in sheet.iter_rows(min_row=2):
                try:
                    if str(row[4].value) == str(date) and row[5].value:
                        booked_times.append(row[5].value)
                except Exception:
                    logger.debug("Ошибка при чтении строки в get_available_times", exc_info=True)
                    continue
            logger.info(f"Найдено забронированных времён для {date}: {len(booked_times)}")
            wb.close()
    except Exception as e:
        logger.error(f"Ошибка при чтении файла db.xlsx: {e}")
        return []

    available_times = [time for time in fixed_times if time not in booked_times]

    # Фильтрация времени для сегодняшней даты
    if date == datetime.date.today().strftime("%d.%m.%Y"):
        now = datetime.datetime.now()
        # Округление до ближайшего получаса
        minutes = now.minute
        if minutes < 30:
            next_time = now.replace(minute=30, second=0, microsecond=0)
        else:
            next_time = (now + datetime.timedelta(hours=1)).replace(minute=0, second=0, microsecond=0)
        next_half_hour = next_time.strftime("%H:%M")
        available_times = [time for time in available_times if time >= next_half_hour]

    return available_times


def ensure_db(path='db.xlsx'):
    """Убедиться, что файл БД существует; если нет — создать с заголовками.

    Ничего не делает, если файл уже существует.
    """
    try:
        with excel_lock:
            if not os.path.exists(path):
                wb = openpyxl.Workbook()
                sheet = wb.active
                # Добавляем колонку remind_before (минуты до процедуры для напоминания)
                sheet.append(['chat_id', 'fio', 'phone', 'massage_type', 'date', 'time', 'remind_before'])
                wb.save(path)
                wb.close()
                logger.info(f"Создан файл базы данных: {path}")
    except Exception as e:
        logger.error(f"Ошибка при создании db.xlsx: {e}")


def get_appointment_by_chat(chat_id):
    """Возвращает запись (fio, phone, massage_type, date, time, row_num) для chat_id или None."""
    try:
        with excel_lock:
            if not os.path.exists('db.xlsx'):
                return None
            wb = openpyxl.load_workbook('db.xlsx')
            sheet = wb.active
            for row in sheet.iter_rows(min_row=2):
                try:
                    if str(row[0].value) == str(chat_id):
                        fio = row[1].value
                        phone = row[2].value
                        massage_type = row[3].value
                        date = row[4].value
                        time_val = row[5].value
                        row_num = row[0].row
                        wb.close()
                        return fio, phone, massage_type, date, time_val, row_num
                except Exception:
                    logger.debug(f"Ошибка при чтении строки в get_appointment_by_chat chat_id={chat_id}", exc_info=True)
                    continue
            wb.close()
    except Exception as e:
        logger.exception(f"Ошибка при получении записи по chat_id={chat_id}: {e}")
    return None


def log_incoming(message):
    """Логирование входящих сообщений: кто, chat_id, текст/контакт и т.д."""
    try:
        user = getattr(message, 'from_user', None)
        if user:
            user_info = f"id={getattr(user, 'id', '')} name={getattr(user,'first_name','')} {getattr(user,'last_name','') or ''} username={getattr(user,'username','') or ''}"
        else:
            user_info = 'unknown'
        text = getattr(message, 'text', None)
        contact = getattr(message, 'contact', None)
        chat_id = getattr(getattr(message, 'chat', None), 'id', None)
        logger.info(f"IN  chat_id={chat_id} from={user_info} text={text!r} contact={contact}")
    except Exception as e:
        logger.exception(f"Ошибка логирования входящего сообщения: {e}")


def log_outgoing(chat_id, text, **kwargs):
    """Логирование исходящих сообщений от бота."""
    try:
        logger.info(f"OUT chat_id={chat_id} text={text!r} extras={{{', '.join(f'{k}={v}' for k,v in kwargs.items())}}}")
    except Exception as e:
        logger.exception(f"Ошибка логирования исходящего сообщения: {e}")