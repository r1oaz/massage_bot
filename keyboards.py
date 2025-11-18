
from telebot import types
import openpyxl
import logging
from utils import excel_lock

logger = logging.getLogger(__name__)

def main_menu_markup(chat_id=None):
    """Главное меню. Если указан chat_id и он есть в базе, добавляет кнопку 'Действия'."""
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
    markup.add(
        types.KeyboardButton("Прайс"),
        types.KeyboardButton("Записаться"),
    )
    # Если передан chat_id, проверим наличие записи в db.xlsx и добавим кнопку 'Действия'
    if chat_id is not None:
        try:
            logger.info(f"Чтение db.xlsx для проверки наличия записи chat_id={chat_id}")
            with excel_lock:
                book = openpyxl.load_workbook('db.xlsx')
                sheet = book.active
                for row in sheet.iter_rows(min_row=2):
                    try:
                        if str(row[0].value) == str(chat_id):
                            markup.add(types.KeyboardButton("Действия"))
                            logger.info(f"Найдена запись для chat_id={chat_id}, добавляем кнопку 'Действия'")
                            break
                    except Exception:
                        logger.debug("Ошибка при проверке строки в main_menu_markup", exc_info=True)
                        continue
                book.close()
        except Exception:
            # Если не удалось прочитать БД, не мешаем пользователю — просто не добавляем кнопку
            pass
    return markup

def massage_types_markup():
    """Клавиатура с типами массажа из price.xlsx"""
    markup = types.ReplyKeyboardMarkup(resize_keyboard=True)
    try:
        from utils import excel_lock
        with excel_lock:
            book = openpyxl.load_workbook("price.xlsx")
            sheet = book.active
            for row in sheet.iter_rows(min_row=2, max_col=1):
                if row[0].value:
                    markup.add(str(row[0].value))
            book.close()
    except Exception:
        markup.add("Ошибка загрузки типов массажа")
    return markup

