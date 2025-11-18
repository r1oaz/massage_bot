import os
import openpyxl
import datetime
import types as _types

from types import SimpleNamespace
try:
    import telebot
    # telebot.types doesn't include SimpleNamespace; add it for our test harness
    telebot.types.SimpleNamespace = _types.SimpleNamespace
except Exception:
    # if telebot isn't available for some reason, continue; other parts of the test mock bot
    pass

ROOT = os.path.dirname(os.path.dirname(__file__))
DB_PATH = os.path.join(ROOT, 'db.xlsx')
PRICE_PATH = os.path.join(ROOT, 'price.xlsx')

def ensure_files():
    # price.xlsx with two massage types
    if not os.path.exists(PRICE_PATH):
        wb = openpyxl.Workbook()
        s = wb.active
        s.append(['Type', 'Price'])
        s.append(['Relax', '1000'])
        s.append(['Therapeutic', '1500'])
        wb.save(PRICE_PATH)
        wb.close()
    # db.xlsx with headers
    if not os.path.exists(DB_PATH):
        wb = openpyxl.Workbook()
        s = wb.active
        s.append(['chat_id', 'fio', 'phone', 'massage_type', 'date', 'time'])
        wb.save(DB_PATH)
        wb.close()

class FakeContact:
    def __init__(self, phone_number):
        self.phone_number = phone_number

class FakeChat:
    def __init__(self, id):
        self.id = id

class FakeMessage:
    def __init__(self, chat_id, text=None, contact=None):
        self.chat = SimpleNamespace(id=chat_id)
        self.text = text
        self.contact = contact

def monkeypatch_bot(bot):
    # Replace bot.send_message and registration functions with stubs that print
    def send_message(chat_id, text, **kwargs):
        print(f"BOT SEND -> chat_id={chat_id}: {text}")
        return None

    def register_next_step_handler(msg, fn, *args, **kwargs):
        print(f"BOT REGISTER next handler: {fn.__name__} args={args}")

    def register_next_step_handler_by_chat_id(chat_id, fn, *args, **kwargs):
        print(f"BOT REGISTER_BY_CHAT next handler for {chat_id}: {fn.__name__} args={args}")

    bot.send_message = send_message
    bot.register_next_step_handler = register_next_step_handler
    bot.register_next_step_handler_by_chat_id = register_next_step_handler_by_chat_id

def print_db():
    wb = openpyxl.load_workbook(DB_PATH)
    s = wb.active
    print('Current DB rows:')
    for row in s.iter_rows(values_only=True):
        print(row)
    wb.close()

def run():
    ensure_files()
    # import app modules
    from config import bot
    import reg_mas
    import reminders

    monkeypatch_bot(bot)

    # Simulate registration flow
    chat_id = 12345
    print('\n--- Start registration flow ---')
    # start -> ask fio
    msg = FakeMessage(chat_id, text='start')
    reg_mas.register_massage(msg)

    # User enters FIO
    msg_fio = FakeMessage(chat_id, text='Ivan Petrov')
    reg_mas.ask_phone_number(msg_fio)

    # User sends contact
    contact = FakeContact('+79991234567')
    msg_contact = FakeMessage(chat_id, contact=contact)
    reg_mas.handle_phone_number(msg_contact, 'Ivan Petrov')

    # User selects massage type
    # call ask_date_step directly
    msg_type = FakeMessage(chat_id, text='Relax')
    reg_mas.ask_date_step(msg_type, 'Ivan Petrov', '+79991234567')

    # Instead of relying on register_next_step_handler, simulate selecting a date
    from utils import get_available_times
    today = datetime.date.today().strftime('%d.%m.%Y')
    available = get_available_times(today)
    if not available:
        print('No available times for today, test will pick tomorrow if possible')
        tomorrow = (datetime.date.today() + datetime.timedelta(days=1)).strftime('%d.%m.%Y')
        chosen_date = tomorrow
        available = get_available_times(chosen_date)
    else:
        chosen_date = today

    chosen_time = available[0]
    print(f'Chosen date/time: {chosen_date} {chosen_time}')

    # Simulate confirmation flow: user selects time then confirms
    msg_time = FakeMessage(chat_id, text=chosen_time)
    # call confirm_registration directly
    reg_mas.confirm_registration(msg_time, 'Ivan Petrov', '+79991234567', 'Relax', chosen_date)

    # user confirms
    msg_confirm = FakeMessage(chat_id, text='Да')
    reg_mas.save_registration(msg_confirm, 'Ivan Petrov', '+79991234567', 'Relax', chosen_date, chosen_time)

    print_db()

    # Test reminder send and response handling
    print('\n--- Test reminder and response ---')
    reminders.send_reminder(chat_id, chosen_date, chosen_time)
    # simulate user confirming from reminder
    msg_rem = FakeMessage(chat_id, text='Подтвердить посещение')
    reminders.handle_reminder_response(msg_rem, chosen_date, chosen_time)

    # simulate user cancelling from reminder
    reminders.send_reminder(chat_id, chosen_date, chosen_time)
    msg_rem2 = FakeMessage(chat_id, text='Отменить запись')
    reminders.handle_reminder_response(msg_rem2, chosen_date, chosen_time)

    print_db()

if __name__ == '__main__':
    run()
